import os
from docx import Document
from docx.shared import Inches
from docx.enum.text import WD_ALIGN_PARAGRAPH

# ============================================================
# CONFIGURATION
# ============================================================
# List of all directories you want to scan
BASE_DIRS = [
    "training/new environment final/baseline/quantum",
    "training/new environment final/baseline/classical"
]

OUTPUT_FILENAME = "Combined_Training_Report.docx"

# ============================================================
# MAIN SCRIPT
# ============================================================
def generate_docx():
    # 1. Create Document
    doc = Document()
    doc.add_heading('Hybrid vs Classical Training Report', 0)
    
    found_any = False

    # 2. Loop through each base directory (Quantum, Classical)
    for base_dir in BASE_DIRS:
        if not os.path.exists(base_dir):
            print(f"[WARN] Directory not found: {base_dir}")
            continue

        # Get the category name (e.g., "quantum" or "classical") from the path
        category_name = os.path.basename(base_dir).upper()
        
        print(f"\nScanning Category: {category_name}...")
        
        # Add a Main Section Header to the Doc
        doc.add_page_break()
        section_head = doc.add_heading(f"=== {category_name} MODEL RUNS ===", level=1)
        section_head.alignment = WD_ALIGN_PARAGRAPH.CENTER

        runs = sorted(os.listdir(base_dir))

        for run_name in runs:
            run_path = os.path.join(base_dir, run_name)
            plot_dir = os.path.join(run_path, "results", "plots")

            # Check validity
            if not os.path.isdir(run_path) or "dont use" in run_name.lower():
                continue
            
            if not os.path.exists(plot_dir):
                # print(f"   [SKIP] No plots in {run_name}")
                continue

            print(f"   [ADD]  {run_name}")
            found_any = True

            # --- Add Run Sub-Header ---
            # doc.add_page_break() # Optional: Uncomment to force new page per run
            run_head = doc.add_heading(run_name, level=2)
            run_head.alignment = WD_ALIGN_PARAGRAPH.LEFT

            # --- Add Images ---
            images = sorted([f for f in os.listdir(plot_dir) if f.endswith(".png")])

            if not images:
                doc.add_paragraph("(No PNG images found)")
                continue

            # Add images in a grid or sequence
            for img_name in images:
                img_path = os.path.join(plot_dir, img_name)
                
                # Clean Label
                label = img_name.replace(".png", "").replace("_", " ").title()
                
                # Add label text
                p = doc.add_paragraph()
                p.add_run(label).bold = True
                p.alignment = WD_ALIGN_PARAGRAPH.CENTER

                # Add Picture
                try:
                    doc.add_picture(img_path, width=Inches(5.5))
                    last_p = doc.paragraphs[-1] 
                    last_p.alignment = WD_ALIGN_PARAGRAPH.CENTER
                    doc.add_paragraph("\n") # Spacing
                except Exception as e:
                    print(f"      [ERR] Failed to add {img_name}: {e}")

    # 3. Save
    if found_any:
        doc.save(OUTPUT_FILENAME)
        print(f"\n✅ Report generated: {OUTPUT_FILENAME}")
    else:
        print("\n❌ No valid plot folders found in any specified directory.")

if __name__ == "__main__":
    generate_docx()
