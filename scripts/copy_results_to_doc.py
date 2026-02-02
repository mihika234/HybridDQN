import os
from docx import Document
from docx.shared import Inches
from docx.enum.text import WD_ALIGN_PARAGRAPH

BASE_DIR = os.path.expanduser("~/HybridDQN/training")
OUTPUT_DOC = os.path.expanduser("~/HybridDQN/training_results_summary.docx")

IMAGE_EXTS = (".png", ".jpg", ".jpeg")

# -------------------------------------------------
# Helpers
# -------------------------------------------------
def is_classical(folder):
    name = folder.lower()
    return name.startswith("classical") or name.startswith("class ")

def is_quantum(folder):
    name = folder.lower()
    return name.startswith("quantum") or name.startswith("quan")

def classify_env(folder_name):
    name = folder_name.lower()
    if name.endswith("new new"):
        return "nonlinear"
    elif name.endswith("new"):
        return "intermediate"
    else:
        return "baseline"

def env_title(env_key):
    return {
        "baseline": "Environment A: Baseline Linear MEC",
        "intermediate": "Environment B: Congestion-Aware MEC",
        "nonlinear": "Environment C: Fully Coupled Non-Linear MEC",
    }[env_key]

def env_description(env_key):
    return {
        "baseline": (
            "This environment uses a linear MEC model with independent fog nodes, "
            "linear energy consumption, and hard deadline-based task dropping."
        ),
        "intermediate": (
            "This environment introduces congestion-awareness through a hockey-stick "
            "fog capacity penalty and probabilistic soft drops, while retaining a linear "
            "energy model and no inter-fog coupling."
        ),
        "nonlinear": (
            "This environment represents the proposed fully coupled MEC system, "
            "incorporating fog-to-fog spillover coupling, hockey-stick congestion effects, "
            "convex non-linear energy consumption, and probabilistic soft drops."
        ),
    }[env_key]

def extract_params(folder_name):
    return folder_name.replace("classical", "") \
                      .replace("quantum", "") \
                      .replace("class", "") \
                      .replace("quan", "") \
                      .replace("new new", "") \
                      .replace("new", "") \
                      .strip()

def find_plot_dirs(run_dir):
    candidates = [
        os.path.join(run_dir, "plots"),
        os.path.join(run_dir, "results", "plots"),
    ]
    return [d for d in candidates if os.path.isdir(d)]

# -------------------------------------------------
# Document creation
# -------------------------------------------------
doc = Document()
doc.add_heading("Training Results Summary", level=0)

for env_key in ["baseline", "intermediate", "nonlinear"]:
    doc.add_heading(env_title(env_key), level=1)
    doc.add_paragraph(env_description(env_key))

    for mode in ["classical", "quantum"]:
        doc.add_heading(mode.capitalize(), level=2)

        mode_dir = os.path.join(BASE_DIR, mode)
        if not os.path.isdir(mode_dir):
            doc.add_paragraph(f"No {mode} experiments found.")
            continue

        found_any = False

        for run in sorted(os.listdir(mode_dir)):
            run_path = os.path.join(mode_dir, run)
            if not os.path.isdir(run_path):
                continue

            if mode == "classical" and not is_classical(run):
                continue
            if mode == "quantum" and not is_quantum(run):
                continue
            if classify_env(run) != env_key:
                continue

            plot_dirs = find_plot_dirs(run_path)
            if not plot_dirs:
                continue

            found_any = True
            params = extract_params(run)

            doc.add_heading(f"Run: {run}", level=3)
            doc.add_paragraph(f"Parameters: {params}")

            for pdir in plot_dirs:
                for img in sorted(os.listdir(pdir)):
                    if img.lower().endswith(IMAGE_EXTS):
                        img_path = os.path.join(pdir, img)
                        doc.add_picture(img_path, width=Inches(5.5))
                        cap = doc.add_paragraph(img.replace("_", " ").replace(".png", ""))
                        cap.alignment = WD_ALIGN_PARAGRAPH.CENTER

        if not found_any:
            doc.add_paragraph(f"No runs found for this environment.")

doc.save(OUTPUT_DOC)
print("✅ Document created at:", OUTPUT_DOC)

